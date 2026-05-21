from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin   = Cm(2.0)
    section.right_margin  = Cm(2.0)

# ── Color palette ─────────────────────────────────────────────────────────────
BLACK     = RGBColor(0x0D, 0x0D, 0x0D)
DARK_GRAY = RGBColor(0x3A, 0x3A, 0x3A)
MID_GRAY  = RGBColor(0x6B, 0x6B, 0x6B)
LIGHT_BG  = RGBColor(0xF5, 0xF5, 0xF5)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT    = RGBColor(0x1A, 0x1A, 0x1A)   # near-black accent
RULE_CLR  = RGBColor(0xE0, 0xE0, 0xE0)

# ── Helpers ───────────────────────────────────────────────────────────────────
def set_cell_bg(cell, rgb: RGBColor):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  f'{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}')
    tcPr.append(shd)

def set_cell_border(cell, top=None, bottom=None, left=None, right=None):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side, val in [('top',top),('bottom',bottom),('left',left),('right',right)]:
        if val:
            el = OxmlElement(f'w:{side}')
            el.set(qn('w:val'),   val.get('val','single'))
            el.set(qn('w:sz'),    val.get('sz','4'))
            el.set(qn('w:space'),'0')
            el.set(qn('w:color'), val.get('color','000000'))
            tcBorders.append(el)
    tcPr.append(tcBorders)

def no_border(cell):
    for side in ('top','bottom','left','right'):
        set_cell_border(cell, **{side: {'val':'nil','sz':'0','color':'FFFFFF'}})

def add_run(para, text, bold=False, size=10, color=BLACK, italic=False):
    run = para.add_run(text)
    run.bold   = bold
    run.italic = italic
    run.font.size  = Pt(size)
    run.font.color.rgb = color
    run.font.name  = 'Calibri'
    return run

def heading(text, size=22, color=BLACK, space_before=6, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.name = 'Calibri'
    return p

def subheading(text, size=11, color=MID_GRAY):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(14)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.name = 'Calibri'
    return p

def section_label(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after  = Pt(6)
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(8)
    r.font.color.rgb = MID_GRAY
    r.font.name = 'Calibri'
    # letter spacing via XML
    rPr = r._r.get_or_add_rPr()
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:val'), '120')
    rPr.append(spacing)
    return p

def divider():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    pPr  = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bot  = OxmlElement('w:bottom')
    bot.set(qn('w:val'),   'single')
    bot.set(qn('w:sz'),    '4')
    bot.set(qn('w:space'), '1')
    bot.set(qn('w:color'), 'E0E0E0')
    pBdr.append(bot)
    pPr.append(pBdr)
    return p

def body(text, color=DARK_GRAY, size=9.5, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(space_after)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.name = 'Calibri'
    return p

def note_box(text):
    """Light-bg callout block."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = tbl.cell(0, 0)
    set_cell_bg(cell, LIGHT_BG)
    no_border(cell)
    cell.width = Inches(6.3)
    para = cell.paragraphs[0]
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after  = Pt(6)
    para.paragraph_format.left_indent  = Pt(8)
    para.paragraph_format.right_indent = Pt(8)
    r = para.add_run(text)
    r.font.size = Pt(9)
    r.font.color.rgb = DARK_GRAY
    r.font.name = 'Calibri'
    r.italic = True
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)

# ═══════════════════════════════════════════════════════════════════════════════
# HEADER
# ═══════════════════════════════════════════════════════════════════════════════
heading('ICODA × Knoah', size=28, space_before=0, space_after=2)
subheading('Global Launch Proposal  ·  60-Day Sprint  ·  June – August 2026', size=11)
divider()

# ═══════════════════════════════════════════════════════════════════════════════
# INVESTMENT SUMMARY
# ═══════════════════════════════════════════════════════════════════════════════
section_label('Investment Summary')

cols  = ['', 'PILOT', 'GROWTH', 'FULL SCALE']
rows  = [
    ('Markets',                     '2',         '4',              '4'),
    ('',                            'Russia + Turkey', 'All four', 'All four'),
    ('ICODA service fee',           '$24K / month',  '$44K / month',  '$90K / month'),
    ('Total · 60 days',             '$48K',      '$88K',           '$180K'),
    ('Media budget (client-managed)','—',         '+$10K / month',  '+$26K / month'),
    ('Service fee per market / mo', '$12K',       '$11K',           '$22.5K'),
]

HDR_BG  = RGBColor(0x0D, 0x0D, 0x0D)
ROW_ALT = RGBColor(0xF9, 0xF9, 0xF9)

t = doc.add_table(rows=1 + len(rows), cols=4)
t.style = 'Table Grid'
t.alignment = WD_TABLE_ALIGNMENT.LEFT

# col widths
widths = [Cm(5.2), Cm(3.6), Cm(3.6), Cm(3.8)]
for i, row in enumerate(t.rows):
    for j, cell in enumerate(row.cells):
        cell.width = widths[j]

# header row
hdr = t.rows[0]
for j, label in enumerate(cols):
    cell = hdr.cells[j]
    set_cell_bg(cell, HDR_BG)
    no_border(cell)
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_before = Pt(5)
    para.paragraph_format.space_after  = Pt(5)
    r = para.add_run(label)
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = WHITE
    r.font.name = 'Calibri'

# data rows
for i, (label, p, g, fs) in enumerate(rows):
    row  = t.rows[i + 1]
    bg   = ROW_ALT if i % 2 == 0 else WHITE
    vals = [label, p, g, fs]
    for j, (cell, val) in enumerate(zip(row.cells, vals)):
        set_cell_bg(cell, bg)
        no_border(cell)
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.space_before = Pt(4)
        para.paragraph_format.space_after  = Pt(4)
        is_price = 'month' in val or val.startswith('$') or val.startswith('+$')
        r = para.add_run(val)
        r.bold = (label in ('ICODA service fee','Total · 60 days') and j > 0) or (j == 0 and label not in ('',''))
        r.font.size = Pt(9 if not is_price else 9.5)
        r.font.color.rgb = BLACK if is_price else DARK_GRAY
        r.font.name = 'Calibri'

doc.add_paragraph().paragraph_format.space_after = Pt(2)
note_box(
    'Why Pilot costs more per market than Growth: fixed infrastructure (account management, '
    'community setup, briefing cycles) is the same for 2 or 4 markets — spread over fewer '
    'markets, the per-unit cost is higher. Standard agency economics for boutique vs. scale engagements.'
)

# ═══════════════════════════════════════════════════════════════════════════════
# DELIVERABLES — KOLs
# ═══════════════════════════════════════════════════════════════════════════════
section_label('Deliverables · KOLs')

kol_rows = [
    ('Tier',                   'Mid-tier',      'Tier 2 + Tier 3 mix',       'Top-tier incl. Tier 1'),
    ('Followers range',        '50K — 200K',    '50K — 500K',                '200K — 1M'),
    ('KOLs per market',        '4 — 6',         '6 — 10',                    '10 — 12'),
    ('Total KOLs (all markets)','8 — 12',        '24 — 40',                   '40 — 48'),
    ('Combined audience reach','500K+',          '2M+',                       '5M+'),
    ('Launch coordination',    'Rolling',        'Rolling',                   'Synchronized blast'),
]

def mini_table(header_cols, data_rows, widths_cm=None):
    if widths_cm is None:
        widths_cm = [5.2, 3.6, 3.6, 3.8]
    t = doc.add_table(rows=1 + len(data_rows), cols=len(header_cols))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, row in enumerate(t.rows):
        for j, cell in enumerate(row.cells):
            cell.width = Cm(widths_cm[j])
    hdr = t.rows[0]
    for j, label in enumerate(header_cols):
        cell = hdr.cells[j]
        set_cell_bg(cell, HDR_BG)
        no_border(cell)
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.space_before = Pt(4)
        para.paragraph_format.space_after  = Pt(4)
        r = para.add_run(label)
        r.bold = True; r.font.size = Pt(9); r.font.color.rgb = WHITE; r.font.name = 'Calibri'
    for i, dr in enumerate(data_rows):
        row = t.rows[i + 1]
        bg  = ROW_ALT if i % 2 == 0 else WHITE
        for j, (cell, val) in enumerate(zip(row.cells, dr)):
            set_cell_bg(cell, bg)
            no_border(cell)
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.space_before = Pt(4)
            para.paragraph_format.space_after  = Pt(4)
            r = para.add_run(val)
            r.font.size = Pt(9); r.font.color.rgb = DARK_GRAY; r.font.name = 'Calibri'
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

mini_table(['', 'PILOT', 'GROWTH', 'FULL SCALE'], kol_rows)

# Community
section_label('Deliverables · Community')
mini_table(['', 'PILOT', 'GROWTH', 'FULL SCALE'], [
    ('Platforms',                 'Telegram',         'Telegram + Facebook',           'Telegram + Facebook'),
    ('Community setup',           '✓',                '✓',                             '✓'),
    ('Local-language management', '30 days',          'Daily · 60 days',               'Daily · 60 days + strategy'),
    ('Competition mechanic',      'Via KOL channels', 'Active seeding + optimization', 'Active seeding + optimization'),
])

# PR
section_label('Deliverables · PR')
mini_table(['', 'PILOT', 'GROWTH', 'FULL SCALE'], [
    ('Placements per market',         '2',                ' 3',                    '5+'),
    ('Coverage type',                 'Local crypto',     'Local + 1 international','Local + international'),
    ('Investment announcement blast', '—',                'Coordinated',           '20+ outlets · launch day'),
    ('Total guaranteed placements',   '4',                '12',                    '25+'),
])

# Paid Media
section_label('Deliverables · Paid Media')
mini_table(['', 'PILOT', 'GROWTH', 'FULL SCALE'], [
    ('Ad spend (client-managed)',    '—',         '$2K — 3K / market',   '$5K — 8K / market'),
    ('Platforms',                    '—',         'Meta + Google',        'Meta + Google + X'),
    ('Campaign management',          '—',         '✓',                    '✓ premium'),
    ('Competitive intelligence',     '—',         '—',                    'Real-time monitoring'),
])

# Creative
section_label('Deliverables · Creative')
mini_table(['', 'PILOT', 'GROWTH', 'FULL SCALE'], [
    ('Asset production',       'Shared (adapted per language)', 'Per-market localization',    'Per-market · premium'),
    ('Value prop',             'Language only',                  'Full per-market hooks',      'Full per-market · no shared assets'),
])

# Reporting
section_label('Deliverables · Reporting & Management')
mini_table(['', 'PILOT', 'GROWTH', 'FULL SCALE'], [
    ('Reporting cadence',    'Monthly',  'Bi-weekly',            'Weekly strategy calls'),
    ('Account management',   'Shared',   'Shared',               'Dedicated'),
    ('Attribution model',    'Standard', 'Standard + channel split', 'Full attribution stack'),
])

# ═══════════════════════════════════════════════════════════════════════════════
# EXPECTED RESULTS
# ═══════════════════════════════════════════════════════════════════════════════
section_label('Expected Results · 60 Days')

res_rows = [
    ('New registered users',            '1K — 3K',    '8K — 20K',   '30K — 60K'),
    ('Community members',               '2K — 5K',    '10K+',       '25K+'),
    ('Media placements',                '4',          '12+',        '25+'),
    ('KOL impressions',                 '500K+',      '2M+',        '5M+'),
    ('Est. cost per user (service fee)', '$16 — 48',   '$4.4 — 11',  '$3 — 6'),
]

t2 = doc.add_table(rows=1 + len(res_rows), cols=4)
t2.style = 'Table Grid'
t2.alignment = WD_TABLE_ALIGNMENT.LEFT
widths2 = [Cm(5.2), Cm(3.6), Cm(3.6), Cm(3.8)]
for i, row in enumerate(t2.rows):
    for j, cell in enumerate(row.cells):
        cell.width = widths2[j]

hdr2 = t2.rows[0]
for j, label in enumerate(['', 'PILOT', 'GROWTH', 'FULL SCALE']):
    cell = hdr2.cells[j]
    set_cell_bg(cell, HDR_BG)
    no_border(cell)
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after  = Pt(4)
    r = para.add_run(label)
    r.bold = True; r.font.size = Pt(9); r.font.color.rgb = WHITE; r.font.name = 'Calibri'

HIGHLIGHT_ROWS = {0, 4}
for i, dr in enumerate(res_rows):
    row = t2.rows[i + 1]
    bg  = ROW_ALT if i % 2 == 0 else WHITE
    for j, (cell, val) in enumerate(zip(row.cells, dr)):
        set_cell_bg(cell, bg)
        no_border(cell)
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.space_before = Pt(4)
        para.paragraph_format.space_after  = Pt(4)
        r = para.add_run(val)
        r.bold = i in HIGHLIGHT_ROWS
        r.font.size = Pt(9.5 if i in HIGHLIGHT_ROWS else 9)
        r.font.color.rgb = BLACK if i in HIGHLIGHT_ROWS else DARK_GRAY
        r.font.name = 'Calibri'

doc.add_paragraph().paragraph_format.space_after = Pt(4)

# ═══════════════════════════════════════════════════════════════════════════════
# WHEN TO CHOOSE
# ═══════════════════════════════════════════════════════════════════════════════
section_label('When to Choose')

wtc_rows = [
    ('Best for',     'Risk-controlled test before full rollout',
                     '4-market launch sprint — recommended default for July window',
                     'Maximum amplification for the investment announcement'),
    ('Ideal if',     'Board approval pending; needs proof-of-fit data first',
                     'Budget confirmed; July deadline is fixed',
                     'Announcement-day impact is the primary KPI'),
    ('Trade-off',    'Misses BR + VN in Wave 1; partial announcement coverage',
                     'Requires $10K/mo client-managed media budget',
                     'Highest total investment; top-tier KOL slots must be confirmed'),
    ('Path forward', 'Upgradeable to Growth at Month 2 on Wave 1 data', '—', '—'),
]
mini_table(['', 'PILOT', 'GROWTH', 'FULL SCALE'], wtc_rows, widths_cm=[3.8, 4.4, 4.4, 3.6])

# ═══════════════════════════════════════════════════════════════════════════════
# MEDIA BUDGET NOTE
# ═══════════════════════════════════════════════════════════════════════════════
section_label('Media Budget · How It Works')
note_box(
    'The media budget (Growth: $10K/mo · Full Scale: $26K/mo) is client-owned spend paid directly '
    'to Meta, Google, and X. ICODA sets up, manages, and optimizes the campaigns — you see every dollar '
    'in your own ad account. Separated from the service fee intentionally: full visibility, full control.'
)

# ═══════════════════════════════════════════════════════════════════════════════
# PAYMENT STRUCTURE
# ═══════════════════════════════════════════════════════════════════════════════
section_label('Payment Structure')
mini_table(['', 'Month 1', 'Month 2', 'Total'], [
    ('Pilot',      '$24K', '$24K', '$48K'),
    ('Growth',     '$44K', '$44K', '$88K'),
    ('Full Scale', '$90K', '$90K', '$180K'),
], widths_cm=[4.0, 4.0, 4.0, 4.2])

doc.add_paragraph().paragraph_format.space_after = Pt(2)
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after  = Pt(2)
add_run(p, 'Month 1', bold=True, size=9, color=DARK_GRAY)
add_run(p, ' — campaign build + Wave 1 launch (Russia + Turkey).  ', size=9, color=MID_GRAY)
add_run(p, 'Month 2', bold=True, size=9, color=DARK_GRAY)
add_run(p, ' — Wave 2 activation (Brazil + Vietnam) + investment-announcement blast.', size=9, color=MID_GRAY)

p2 = doc.add_paragraph()
p2.paragraph_format.space_before = Pt(2)
p2.paragraph_format.space_after  = Pt(4)
add_run(p2, 'No ongoing retainer — the engagement has a defined endpoint at Week 8.', size=9, color=MID_GRAY, italic=True)

# ═══════════════════════════════════════════════════════════════════════════════
# TIMELINE
# ═══════════════════════════════════════════════════════════════════════════════
section_label('Timeline')

timeline = [
    ('Week 1',       'Contract signed · KOL outreach begins · communities set up (RU + TR)'),
    ('Week 2',       'KOL posts live · Telegram seeding · paid campaigns activated'),
    ('Week 3',       'Wave 1 live (RU + TR) · first attribution data · Wave 2 prep begins'),
    ('Week 4',       'Creative re-tuned for BR + VN using Week 3 data'),
    ('Week 5',       'Wave 2 live (BR + VN) · all four markets active'),
    ('Weeks 6 — 8',  'Investment-announcement amplification blast across all markets'),
    ('End Month 2',  'Full performance report + Month 3 optimization plan'),
]

t3 = doc.add_table(rows=len(timeline), cols=2)
t3.style = 'Table Grid'
t3.alignment = WD_TABLE_ALIGNMENT.LEFT
for i, row in enumerate(t3.rows):
    row.cells[0].width = Cm(3.2)
    row.cells[1].width = Cm(13.0)
    bg = ROW_ALT if i % 2 == 0 else WHITE
    label, desc = timeline[i]
    for j, (cell, val) in enumerate(zip(row.cells, [label, desc])):
        set_cell_bg(cell, bg)
        no_border(cell)
        para = cell.paragraphs[0]
        para.paragraph_format.space_before = Pt(5)
        para.paragraph_format.space_after  = Pt(5)
        r = para.add_run(val)
        r.bold = (j == 0)
        r.font.size = Pt(9)
        r.font.color.rgb = BLACK if j == 0 else DARK_GRAY
        r.font.name = 'Calibri'

doc.add_paragraph().paragraph_format.space_after = Pt(4)
note_box('Critical path: contract must be signed by June 1 to make the July announcement window. '
         'KOL briefing and community setup requires a full two weeks — there is no shortcut.')

# ═══════════════════════════════════════════════════════════════════════════════
# FOOTER
# ═══════════════════════════════════════════════════════════════════════════════
divider()
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after  = Pt(0)
add_run(p, 'Ruslan Pivnev  ·  CBDO, ICODA  ·  ', size=8.5, color=MID_GRAY)
add_run(p, 'rp@icoda.io', size=8.5, color=MID_GRAY)
add_run(p, '  ·  @ruslan_bsag  ·  icoda.io', size=8.5, color=MID_GRAY)

doc.save('/home/user/testcoldoutreach/knoah-package-comparison.docx')
print("Done.")
