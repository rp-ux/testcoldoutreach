"""
Читает сообщения, голосовые и файлы из указанных чатов.
Требования: pip3 install telethon python-dotenv openai-whisper pdfplumber python-docx openpyxl
"""
import asyncio
import json
import os
import tempfile
from datetime import datetime
from dotenv import load_dotenv
from telethon import TelegramClient
from telethon.tl.types import (
    MessageMediaDocument, MessageMediaPhoto,
    DocumentAttributeAudio, DocumentAttributeFilename
)

load_dotenv()

API_ID = int(os.environ["TELEGRAM_API_ID"])
API_HASH = os.environ["TELEGRAM_API_HASH"]
SESSION = os.environ.get("SESSION_NAME", "ruslan_bsag")

CHATS_TO_READ = ["iTor_0x", "Ivan_ICODA", "avicoda"]
KEYWORDS = ["kaching", "качинг", "kach"]
MESSAGES_LIMIT = 100

# Whisper для транскрипции голосовых
try:
    import whisper
    whisper_model = whisper.load_model("base")
    WHISPER_AVAILABLE = True
    print("Whisper loaded.")
except Exception:
    WHISPER_AVAILABLE = False
    print("Whisper unavailable. Voice messages will be skipped.")


def read_file(path: str, filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()
    try:
        if ext == ".pdf":
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                return "\n".join(p.extract_text() or "" for p in pdf.pages)
        elif ext in (".docx", ".doc"):
            from docx import Document
            doc = Document(path)
            return "\n".join(p.text for p in doc.paragraphs)
        elif ext in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True)
            out = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                out.append(f"[Sheet: {sheet}]")
                for row in ws.iter_rows(values_only=True):
                    line = "\t".join(str(c) if c is not None else "" for c in row)
                    if line.strip():
                        out.append(line)
            return "\n".join(out)
        elif ext in (".txt", ".csv", ".md"):
            with open(path, encoding="utf-8", errors="ignore") as f:
                return f.read()
    except ImportError as e:
        return f"[Библиотека не установлена: {e}]"
    except Exception as e:
        return f"[Ошибка чтения файла: {e}]"
    return f"[Формат {ext} не поддерживается]"


async def transcribe_voice(client, message) -> str:
    if not WHISPER_AVAILABLE:
        return "[Голосовое — Whisper не установлен]"
    with tempfile.NamedTemporaryFile(suffix=".ogg", delete=False) as tmp:
        tmp_path = tmp.name
    try:
        await client.download_media(message, file=tmp_path)
        result = whisper_model.transcribe(tmp_path)
        return f"[Голосовое]: {result['text'].strip()}"
    except Exception as e:
        return f"[Ошибка транскрипции: {e}]"
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


async def process_message(client, msg) -> dict | None:
    if not msg.text and not msg.media:
        return None

    sender = ""
    if msg.sender:
        if hasattr(msg.sender, "username") and msg.sender.username:
            sender = f"@{msg.sender.username}"
        elif hasattr(msg.sender, "first_name"):
            sender = msg.sender.first_name or "Unknown"

    date = msg.date.strftime("%Y-%m-%d %H:%M")
    text = msg.text or ""

    # Голосовое сообщение
    if msg.media and isinstance(msg.media, MessageMediaDocument):
        doc = msg.media.document
        attrs = {type(a).__name__: a for a in doc.attributes}

        if "DocumentAttributeAudio" in attrs:
            audio_attr = attrs["DocumentAttributeAudio"]
            if getattr(audio_attr, "voice", False):
                text = await transcribe_voice(client, msg)
            else:
                text = f"[Аудио файл]"

        elif "DocumentAttributeFilename" in attrs:
            filename = attrs["DocumentAttributeFilename"].file_name
            with tempfile.NamedTemporaryFile(suffix=os.path.splitext(filename)[1], delete=False) as tmp:
                tmp_path = tmp.name
            try:
                await client.download_media(msg, file=tmp_path)
                content = read_file(tmp_path, filename)
                text = f"[Файл: {filename}]\n{content}"
            finally:
                if os.path.exists(tmp_path):
                    os.remove(tmp_path)

    if not text:
        return None

    return {"date": date, "sender": sender, "text": text}


async def get_chat_messages(client, chat_id, limit=100):
    messages = []
    try:
        entity = await client.get_entity(chat_id)
        async for msg in client.iter_messages(entity, limit=limit):
            item = await process_message(client, msg)
            if item:
                messages.append(item)
    except Exception as e:
        print(f"  Ошибка при чтении {chat_id}: {e}")
    return messages


async def find_kaching_dialog(client):
    print("Ищу чат Kaching <> ICODA...")
    async for dialog in client.iter_dialogs():
        name = dialog.name or ""
        if any(k in name.lower() for k in KEYWORDS):
            print(f"  Найден: {name}")
            messages = []
            async for msg in client.iter_messages(dialog.entity, limit=MESSAGES_LIMIT):
                item = await process_message(client, msg)
                if item:
                    messages.append(item)
            return name, messages
    return None, []


async def main():
    client = TelegramClient(SESSION, API_ID, API_HASH)
    await client.start()

    all_data = {}

    name, messages = await find_kaching_dialog(client)
    if messages:
        all_data[name] = messages
        print(f"  Загружено {len(messages)} сообщений")
    else:
        print("  Чат Kaching не найден")

    for username in CHATS_TO_READ:
        print(f"Читаю @{username}...")
        msgs = await get_chat_messages(client, username, MESSAGES_LIMIT)
        if msgs:
            all_data[f"@{username}"] = msgs
            print(f"  Загружено {len(msgs)} сообщений")

    await client.disconnect()

    with open("chats_export.json", "w", encoding="utf-8") as f:
        json.dump(all_data, f, ensure_ascii=False, indent=2)

    print(f"\nГотово! Сохранено в chats_export.json")

    print("\n" + "=" * 60)
    print("ВСЕ СООБЩЕНИЯ:")
    print("=" * 60)
    for chat, msgs in all_data.items():
        print(f"\n--- {chat} ---")
        for msg in msgs[:30]:
            print(f"[{msg['date']}] {msg['sender']}: {msg['text'][:300]}")


if __name__ == "__main__":
    asyncio.run(main())
